#!/usr/bin/env python3
"""生成空白 Office 文档模板（docx/xlsx/pptx），放入扩展 Resources"""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")
os.makedirs(OUT, exist_ok=True)

# ---------- docx：正文空、默认字体设为中英文混排友好 ----------
d = Document()
style = d.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
d.save(os.path.join(OUT, "blank.docx"))

# ---------- xlsx：单个空白工作表 ----------
wb = Workbook()
wb.active.title = "Sheet1"
wb.save(os.path.join(OUT, "blank.xlsx"))

# ---------- pptx：单页空白幻灯片（16:9） ----------
p = Presentation()
p.slide_width = 12192000
p.slide_height = 6858000
p.slides.add_slide(p.slide_layouts[6])   # 6 = 空白版式
p.save(os.path.join(OUT, "blank.pptx"))

for f in sorted(os.listdir(OUT)):
    path = os.path.join(OUT, f)
    print(f"{f:12} {os.path.getsize(path):>7} bytes")
